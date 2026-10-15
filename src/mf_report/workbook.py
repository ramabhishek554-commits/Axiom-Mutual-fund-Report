"""Create and reopen-check the 11-sheet Excel deliverable."""
from io import BytesIO
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .analytics import company_comparison, concentration_analysis, historical_changes, overlap_matrix, sector_exposure
from .extract import validate_holdings

SHEETS=["Executive Summary","Large Cap Funds","Mid Cap Funds","Small Cap Funds","All Holdings","Company Comparison","Fund Overlap","Sector Exposure","Concentration Analysis","Historical Changes","Sources and Validation"]

def _write_df(ws, df):
    if df is None or df.empty:
        ws.append(["No verified data available"])
        return
    ws.append([str(c) for c in df.columns])
    for row in df.itertuples(index=False,name=None):
        vals=[]
        for v in row:
            try: vals.append(None if pd.isna(v) else (v.item() if hasattr(v,"item") else v))
            except (TypeError,ValueError): vals.append(v)
        ws.append(vals)
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:
        cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="17365D")
        cell.alignment=Alignment(wrap_text=True,vertical="top")
    for col in ws.columns:
        length=max((len(str(c.value or "")) for c in list(col)[:200]),default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width=min(max(length+2,12),42)

def build_workbook(holdings, investment_amount=100000, metadata=None):
    clean, issues=validate_holdings(holdings)
    valid=clean[~clean.validation_status.astype(str).eq("invalid_weight")] if "validation_status" in clean else clean
    wb=Workbook(); summary=wb.active; summary.title="Executive Summary"
    summary.append(["Axiom Mutual Fund Intelligence & Excel Analyzer"])
    summary.append(["Assumed investor amount (INR)",float(investment_amount)])
    summary.append(["Holdings rows supplied",int(len(holdings))])
    summary.append(["Rows requiring review",int(valid.get("validation_status",pd.Series(dtype=str)).astype(str).ne("validated").sum())])
    summary.append(["Important","Verify scheme category, portfolio date, source and extracted rows."])
    summary.append(["Disclaimer","Research tool only; not investment advice or return prediction."])
    summary["A1"].font=Font(bold=True,size=14,color="FFFFFF"); summary["A1"].fill=PatternFill("solid",fgColor="17365D")
    summary.column_dimensions["A"].width=42; summary.column_dimensions["B"].width=90
    if not valid.empty: valid=valid.copy(); valid["estimated_allocation_inr"]=valid.apply(lambda r: float(investment_amount)*float(r["weight_pct"])/100,axis=1)
    frames={
      "Large Cap Funds":valid[valid.category.eq("Large Cap")],"Mid Cap Funds":valid[valid.category.eq("Mid Cap")],
      "Small Cap Funds":valid[valid.category.eq("Small Cap")],"All Holdings":valid,
      "Company Comparison":company_comparison(valid,investment_amount),"Fund Overlap":overlap_matrix(valid),
      "Sector Exposure":sector_exposure(valid),"Concentration Analysis":concentration_analysis(valid),
      "Historical Changes":historical_changes(valid)}
    for name,df in frames.items(): _write_df(wb.create_sheet(name),df)
    pieces=[]
    if metadata is not None and not metadata.empty: pieces.append(metadata)
    if not issues.empty: pieces.append(issues)
    source_cols=[c for c in ["scheme","category","valuation_date","source_url","source_pdf","page","extraction_method","validation_status"] if c in valid]
    if source_cols: pieces.append(valid[source_cols])
    _write_df(wb.create_sheet("Sources and Validation"),pd.concat(pieces,ignore_index=True,sort=False) if pieces else pd.DataFrame())
    stream=BytesIO(); wb.save(stream); return stream.getvalue()

def validate_workbook(data):
    try:
        wb=load_workbook(BytesIO(data),read_only=True)
        missing=sorted(set(SHEETS)-set(wb.sheetnames)); wb.close()
        return [f"Missing workbook sheets: {', '.join(missing)}"] if missing else []
    except Exception as exc: return [f"Workbook could not be reopened: {exc}"]
