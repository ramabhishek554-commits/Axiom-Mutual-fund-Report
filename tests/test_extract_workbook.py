import io
import pandas as pd
from openpyxl import load_workbook
from src.mf_report.extract import normalize_company, validate_holdings
from src.mf_report.workbook import build_workbook, validate_workbook, SHEETS

def test_company_normalization(): assert normalize_company("Example Industries Ltd.")=="example"
def test_invalid_weights_flagged():
    d=pd.DataFrame([{"scheme":"A","category":"Large Cap","company":"Issuer","weight_pct":120}])
    checked,issues=validate_holdings(d)
    assert "invalid_weight" in set(checked.validation_status)
    assert not issues.empty
def test_excel_has_eleven_sheets_and_reopens():
    d=pd.DataFrame([{"scheme":"A","category":"Large Cap","company":"Issuer X","weight_pct":12.5,"sector":"IT",
        "instrument_type":"Equity","valuation_date":"2026-08-31","validation_status":"validated","source_url":"https://example.org/factsheet.pdf"}])
    data=build_workbook(d,100000)
    assert validate_workbook(data)==[]
    wb=load_workbook(io.BytesIO(data),read_only=True)
    assert wb.sheetnames==SHEETS
    assert wb["Executive Summary"]["B2"].value==100000
    wb.close()



def test_validation_checks_totals_and_duplicates_per_valuation_date():
    d = pd.DataFrame([
        {"scheme":"A","category":"Large Cap","company":"Issuer X Ltd","weight_pct":60,"valuation_date":"2026-08-31"},
        {"scheme":"A","category":"Large Cap","company":"Issuer Y Ltd","weight_pct":40,"valuation_date":"2026-08-31"},
        {"scheme":"A","category":"Large Cap","company":"Issuer X Ltd","weight_pct":55,"valuation_date":"2026-09-30"},
        {"scheme":"A","category":"Large Cap","company":"Issuer Y Ltd","weight_pct":45,"valuation_date":"2026-09-30"},
    ])
    checked, issues = validate_holdings(d)
    assert len(checked) == 4
    assert issues.empty


def test_validation_still_flags_duplicate_issuer_within_same_snapshot():
    d = pd.DataFrame([
        {"scheme":"A","category":"Large Cap","company":"Issuer X Ltd","weight_pct":40,"valuation_date":"2026-08-31"},
        {"scheme":"A","category":"Large Cap","company":"Issuer X","weight_pct":30,"valuation_date":"2026-08-31"},
    ])
    _checked, issues = validate_holdings(d)
    assert issues.issue.str.contains("Duplicate normalized issuer").any()
